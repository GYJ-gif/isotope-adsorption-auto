from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SHEET_ALIASES = {
    "吸附数据": ("吸附数据", "ads-des", "ads/des", "簾現方象"),
    "Qst结果": ("Qst结果", "Qst", "Qst潤惚"),
    "IAST结果": ("IAST结果", "IAST", "IAST潤惚"),
}


def find_or_create_sheet(workbook, target_name: str):
    for name in SHEET_ALIASES[target_name]:
        if name in workbook.sheetnames:
            sheet = workbook[name]
            if sheet.title != target_name and target_name not in workbook.sheetnames:
                sheet.title = target_name
            return workbook[target_name] if target_name in workbook.sheetnames else sheet
    return workbook.create_sheet(target_name)


def clear_sheet(sheet) -> None:
    if sheet.max_row:
        sheet.delete_rows(1, sheet.max_row)
    if sheet.max_column:
        sheet.delete_cols(1, sheet.max_column)


def read_csv(path: Path) -> tuple[list[str], list[list[float | str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        rows: list[list[float | str]] = []
        for row in reader:
            converted: list[float | str] = []
            for value in row:
                try:
                    converted.append(float(value))
                except ValueError:
                    converted.append(value)
            rows.append(converted)
    return header, rows


def write_table(sheet, start_row: int, start_col: int, title: str, header: list[str], rows: list[list[float | str]]) -> None:
    sheet.cell(start_row, start_col).value = title
    for index, value in enumerate(header, start=start_col):
        sheet.cell(start_row + 1, index).value = value
    for row_index, row in enumerate(rows, start=start_row + 2):
        for col_index, value in enumerate(row, start=start_col):
            sheet.cell(row_index, col_index).value = value
            if isinstance(value, float):
                sheet.cell(row_index, col_index).number_format = "0.000000"


def format_columns(sheet, max_col: int) -> None:
    for col in range(1, max_col + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 20


def sync_results(workbook_path: Path) -> None:
    sample = workbook_path.stem
    sample_dir = workbook_path.parent
    workbook = load_workbook(workbook_path)

    adsorption_sheet = find_or_create_sheet(workbook, "吸附数据")
    if adsorption_sheet.title != "吸附数据" and "吸附数据" not in workbook.sheetnames:
        adsorption_sheet.title = "吸附数据"

    iast_sheet = find_or_create_sheet(workbook, "IAST结果")
    qst_sheet = find_or_create_sheet(workbook, "Qst结果")
    clear_sheet(iast_sheet)
    clear_sheet(qst_sheet)

    iast_root = sample_dir / "IAST结果"
    for temperature, col in (("77", 1), ("87", 10)):
        csv_path = iast_root / f"{sample}-{temperature}K-IAST_Selectivity.csv"
        if not csv_path.exists():
            raise FileNotFoundError(f"缺少 IAST 结果：{csv_path}")
        header, rows = read_csv(csv_path)
        write_table(iast_sheet, 1, col, f"{temperature}K IAST Selectivity", header, rows)

    qst_root = sample_dir / "Qst结果"
    for gas, col in (("H2", 1), ("D2", 6)):
        csv_path = qst_root / f"{sample}-{gas}-77K-87K-Qst_Qst.csv"
        if not csv_path.exists():
            raise FileNotFoundError(f"缺少 Qst 结果：{csv_path}")
        header, rows = read_csv(csv_path)
        write_table(qst_sheet, 1, col, f"{gas} Qst (77K + 87K)", header, rows)

    format_columns(iast_sheet, 16)
    format_columns(qst_sheet, 8)
    workbook.save(workbook_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Write IAST/Qst CSV results back into the sample workbook.")
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    sync_results(args.workbook.resolve())
    print(f"Synced workbook: {args.workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
