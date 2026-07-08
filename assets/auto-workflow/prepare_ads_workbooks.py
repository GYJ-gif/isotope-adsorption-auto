from __future__ import annotations

import argparse
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


ORDER = (("77", "H2"), ("77", "D2"), ("87", "H2"), ("87", "D2"))
MOLAR_VOLUME_STP_CM3_PER_MMOL = 22.414
SAMPLE_RENAMES = {
    "Ag-beta": "Ag-BeTa",
    "Cu-FMR": "Cu-FER",
}


def parse_source_name(path: Path) -> tuple[str, str, str]:
    stem = path.stem
    match = re.search(r"(?i)(?:(77|87)\s*K\s*-?\s*(H2|D2)|(H2|D2)\s*-?\s*(77|87)\s*K)\b", stem)
    if not match:
        raise ValueError("文件名中未识别到 77K/87K 与 H2/D2")

    sample = stem[: match.start()]
    sample = re.sub(r"[_\-\s]+$", "", sample).strip()
    sample = re.sub(r"\s+", " ", sample)
    sample = SAMPLE_RENAMES.get(sample, sample)
    if not sample:
        raise ValueError("文件名中未识别到样品名")

    temperature = match.group(1) or match.group(4)
    gas = match.group(2) or match.group(3)
    return sample, temperature, gas.upper()


def find_isotherm_sheet(workbook) -> object:
    for sheet in workbook.worksheets:
        title = sheet.title.lower()
        if "等温线" in sheet.title or "p)" in title or "p-p" in title or "sheet2" in title:
            return sheet
    if len(workbook.worksheets) >= 2:
        return workbook.worksheets[1]
    raise ValueError("工作簿中没有可用的第 2 张等温线表")


def find_header_row(sheet) -> tuple[int, int, int]:
    for row in range(1, sheet.max_row + 1):
        values = [sheet.cell(row, col).value for col in range(1, sheet.max_column + 1)]
        pressure_col = None
        loading_col = None
        for index, value in enumerate(values, start=1):
            text = "" if value is None else str(value)
            if "P/" in text and "Pa" in text:
                pressure_col = index
            if "V" in text and ("cm" in text or "cm³" in text) and "g" in text:
                loading_col = index
        if pressure_col and loading_col:
            return row, pressure_col, loading_col
    raise ValueError("未找到 P/(Pa) 与 V/(cm³/g.STP) 表头")


def is_number(value) -> bool:
    if value is None:
        return False
    try:
        float(value)
        return True
    except (TypeError, ValueError):
        return False


def convert_row(pressure_pa: float, loading_cm3_g: float) -> tuple[float, float]:
    return pressure_pa / 1000.0, loading_cm3_g / MOLAR_VOLUME_STP_CM3_PER_MMOL


def extract_isotherm_rows(path: Path) -> dict[str, list[tuple[float, float]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = find_isotherm_sheet(workbook)
        header_row, pressure_col, loading_col = find_header_row(sheet)
        adsorption: list[tuple[float, float]] = []
        desorption: list[tuple[float, float]] = []
        previous_pressure = None
        phase = "adsorption"

        for row in range(header_row + 1, sheet.max_row + 1):
            pressure_value = sheet.cell(row, pressure_col).value
            loading_value = sheet.cell(row, loading_col).value
            if not is_number(pressure_value) or not is_number(loading_value):
                if adsorption or desorption:
                    break
                continue

            pressure_pa = float(pressure_value)
            loading_cm3_g = float(loading_value)
            if pressure_pa < 0 or loading_cm3_g < 0:
                raise ValueError(f"第 {row} 行出现负值")

            if previous_pressure is not None and pressure_pa <= previous_pressure:
                phase = "desorption"

            converted = convert_row(pressure_pa, loading_cm3_g)
            if phase == "adsorption":
                adsorption.append(converted)
            else:
                desorption.append(converted)
            previous_pressure = pressure_pa

        if len(adsorption) < 10:
            raise ValueError(f"吸附段数据点少于 10 个：{len(adsorption)}")
        return {"adsorption": adsorption, "desorption": desorption}
    finally:
        workbook.close()


def write_sample_workbook(sample: str, datasets: dict[tuple[str, str], dict[str, list[tuple[float, float]]]], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "吸附数据"
    workbook.create_sheet("Qst结果")
    workbook.create_sheet("IAST结果")

    for index, (temperature, gas) in enumerate(ORDER):
        start_col = index * 3 + 1
        sheet.cell(1, start_col).value = "Isotherm"
        sheet.cell(3, start_col).value = f"{sample}-{gas}-{temperature}K : 000-000 : Adsorption"
        sheet.cell(4, start_col).value = "Absolute Pressure (kPa)"
        sheet.cell(4, start_col + 1).value = "Quantity Adsorbed (mmol/g)"

        dataset = datasets.get((temperature, gas), {"adsorption": [], "desorption": []})
        adsorption_rows = dataset["adsorption"]
        for row_index, (pressure_kpa, loading_mmol_g) in enumerate(adsorption_rows, start=5):
            sheet.cell(row_index, start_col).value = pressure_kpa
            sheet.cell(row_index, start_col + 1).value = loading_mmol_g

        desorption_header_row = 5 + len(adsorption_rows) + 1
        sheet.cell(desorption_header_row, start_col).value = f"{sample}-{gas}-{temperature}K : 000-000 : Desorption"
        sheet.cell(desorption_header_row + 1, start_col).value = "Absolute Pressure (kPa)"
        sheet.cell(desorption_header_row + 1, start_col + 1).value = "Quantity Adsorbed (mmol/g)"
        for row_index, (pressure_kpa, loading_mmol_g) in enumerate(dataset["desorption"], start=desorption_header_row + 2):
            sheet.cell(row_index, start_col).value = pressure_kpa
            sheet.cell(row_index, start_col + 1).value = loading_mmol_g

        max_row = desorption_header_row + 2 + len(dataset["desorption"])
        for row in range(5, max_row):
            sheet.cell(row, start_col).number_format = "0.00000"
            sheet.cell(row, start_col + 1).number_format = "0.000000"

        sheet.column_dimensions[get_column_letter(start_col)].width = 22
        sheet.column_dimensions[get_column_letter(start_col + 1)].width = 24
        sheet.column_dimensions[get_column_letter(start_col + 2)].width = 4

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def default_output_root(input_root: Path) -> Path:
    today = datetime.now().strftime("%y%m%d")
    return input_root.parent / today


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare adsorption workbooks from raw instrument XLSX files.")
    parser.add_argument("--input-root", required=True, type=Path, help="包含原始 .xlsx 文件的文件夹")
    parser.add_argument("--output-root", type=Path, help="输出日期文件夹；默认在输入文件夹上一级按当天日期命名")
    args = parser.parse_args()

    input_root = args.input_root.resolve()
    output_root = (args.output_root or default_output_root(input_root)).resolve()
    files = sorted(path for path in input_root.glob("*.xlsx") if not path.name.startswith("~$"))
    if not files:
        raise SystemExit(f"未找到原始 .xlsx 文件：{input_root}")

    grouped: dict[str, dict[tuple[str, str], Path]] = defaultdict(dict)
    skipped: list[str] = []
    for path in files:
        try:
            sample, temperature, gas = parse_source_name(path)
            key = (temperature, gas)
            if key in grouped[sample]:
                raise ValueError(f"{sample} 的 {temperature}K-{gas} 数据重复")
            grouped[sample][key] = path
        except Exception as exc:
            skipped.append(f"{path.name}\tSKIP\t{exc}")

    report_lines: list[str] = []
    generated_samples: list[str] = []
    for sample in sorted(grouped):
        datasets: dict[tuple[str, str], dict[str, list[tuple[float, float]]]] = {}
        for key in ORDER:
            source = grouped[sample].get(key)
            if not source:
                continue
            try:
                datasets[key] = extract_isotherm_rows(source)
                ads_count = len(datasets[key]["adsorption"])
                des_count = len(datasets[key]["desorption"])
                report_lines.append(f"{sample}\t{key[0]}K-{key[1]}\tOK\t{source.name}\t{ads_count} adsorption points; {des_count} desorption points")
            except Exception as exc:
                report_lines.append(f"{sample}\t{key[0]}K-{key[1]}\tERROR\t{source.name}\t{exc}")

        if datasets:
            output_path = output_root / sample / f"{sample}.xlsx"
            write_sample_workbook(sample, datasets, output_path)
            generated_samples.append(sample)
            report_lines.append(f"{sample}\tWORKBOOK\tOK\t{output_path}")

    report_dir = output_root / "_prepare_logs"
    report_dir.mkdir(parents=True, exist_ok=True)
    report_path = report_dir / "prepare_ads_workbooks_report.tsv"
    report_path.write_text("\n".join(skipped + report_lines) + "\n", encoding="utf-8")
    samples_path = report_dir / "generated_samples.txt"
    samples_path.write_text("\n".join(generated_samples) + "\n", encoding="utf-8")

    print(f"OutputRoot: {output_root}")
    print(f"Samples: {len(grouped)}")
    print(f"Report: {report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
